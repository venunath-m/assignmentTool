# assignment_generator_app.py
import streamlit as st
from duckduckgo_search import duckduckgo_search
import wikipedia
from docx import Document
from docx.shared import Inches
from fpdf import FPDF
from PIL import Image
from io import BytesIO
import requests
import tempfile
import time
from pathlib import Path
from styles import render_header, info_box, floating_icons

# ---------------------------- Wikipedia setup ----------------------------
wikipedia.set_lang("en")
wikipedia.set_rate_limiting(True)

# ---------------------------- Helper Functions ----------------------------
def fetch_wikipedia_content_with_images(keyword, max_points=10, max_images=3, retries=2):
    suggestions = []
    for _ in range(retries + 1):
        try:
            page = wikipedia.page(keyword)
            content = page.content
            images = [img for img in page.images if img.lower().endswith(('.png', '.jpg', '.jpeg'))][:max_images]

            paragraphs = [p.strip() for p in content.split('\n') if p.strip()]
            intro = paragraphs[0] if paragraphs else "[No intro available]"

            key_points = []
            for para in paragraphs[1:]:
                for s in para.split('. '):
                    s = s.strip()
                    if s and s not in key_points:
                        key_points.append(s)
                    if len(key_points) >= max_points:
                        break
                if len(key_points) >= max_points:
                    break

            return intro, key_points, images, suggestions

        except wikipedia.exceptions.DisambiguationError as e:
            suggestions = e.options[:5]
            try:
                page = wikipedia.page(e.options[0])
                content = page.content
                images = [img for img in page.images if img.lower().endswith(('.png', '.jpg', '.jpeg'))][:max_images]
                paragraphs = [p.strip() for p in content.split('\n') if p.strip()]
                intro = paragraphs[0] if paragraphs else "[No intro available]"

                key_points = []
                for para in paragraphs[1:]:
                    for s in para.split('. '):
                        s = s.strip()
                        if s and s not in key_points:
                            key_points.append(s)
                        if len(key_points) >= max_points:
                            break
                    if len(key_points) >= max_points:
                        break
                return intro, key_points, images, suggestions
            except:
                continue
        except wikipedia.exceptions.PageError:
            suggestions = wikipedia.search(keyword)[:5]
            return None, [], [], suggestions
        except:
            time.sleep(1)
    return None, [], [], suggestions

def fetch_duckduckgo_content(keyword, max_points=10, retries=2):
    for _ in range(retries + 1):
        try:
            results = duckduckgo_search(keyword, max_results=5)
            if not results:
                continue

            combined_text = ""
            for r in results:
                if 'body' in r and r['body']:
                    combined_text += r['body'] + ". "
                elif 'title' in r and r['title']:
                    combined_text += r['title'] + ". "

            sentences = [s.strip() for s in combined_text.split('. ') if len(s.strip()) > 15]

            key_points = []
            seen = set()
            for s in sentences:
                if s not in seen:
                    key_points.append(s)
                    seen.add(s)
                if len(key_points) >= max_points:
                    break

            intro = key_points[0] if key_points else combined_text[:200] + "..."
            return intro, key_points

        except:
            time.sleep(1)
    return None, []

def fetch_and_prepare_image(img_url):
    headers = {
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                      "AppleWebKit/537.36 (KHTML, like Gecko) "
                      "Chrome/120.0.0.0 Safari/537.36"
    }
    try:
        response = requests.get(img_url, headers=headers, timeout=10)
        response.raise_for_status()
        img = Image.open(BytesIO(response.content))

        if img.mode != "RGB":
            img = img.convert("RGB")

        tmp_file = tempfile.NamedTemporaryFile(delete=False, suffix=".jpg")
        img.save(tmp_file.name, format="JPEG")
        tmp_file.close()
        return tmp_file.name
    except Exception as e:
        print(f"Failed to process image {img_url}: {e}")
        return None

# ---------------------------- Assignment Generation ----------------------------
def generate_assignment(keywords, max_points=10, max_images=3):
    combined_content = ""
    all_suggestions = {}
    all_images = {}

    for kw in keywords:
        combined_content += f"=== Topic: {kw} ===\n\n"

        intro, key_points, images, suggestions = fetch_wikipedia_content_with_images(
            kw, max_points=max_points, max_images=max_images
        )
        all_suggestions[kw] = suggestions
        all_images[kw] = images

        if not intro:
            ddg_intro, ddg_points = fetch_duckduckgo_content(kw, max_points=max_points)
            if ddg_intro:
                intro = ddg_intro
                key_points = ddg_points
            else:
                intro = "[Could not fetch information. Please check your internet connection or try again.]"
                key_points = ["[No key points available]"]

        if not key_points:
            key_points = ["[No key points available]"]

        conclusion = f"In summary, {kw} is an important topic in its field."
        key_points_str = "\n- ".join(key_points)

        section = f"""1. Introduction
{intro}

2. Key Points
- {key_points_str}

3. Conclusion
{conclusion}

{'='*50}
"""
        combined_content += section

    return combined_content.strip(), all_suggestions, all_images

# ---------------------------- DOCX & PDF Generation ----------------------------
def generate_docx(content, images_dict, title="Assignment"):
    doc = Document()
    doc.add_heading(title, 0)
    
    topics = content.split("="*50)
    for topic_section in topics:
        if not topic_section.strip():
            continue
        doc.add_paragraph(topic_section.strip())
        lines = topic_section.strip().split("\n")
        topic_name = lines[0].replace("=== Topic: ", "").strip() if lines else None
        if topic_name and topic_name in images_dict:
            for img_url in images_dict[topic_name]:
                tmp_file = fetch_and_prepare_image(img_url)
                if tmp_file:
                    try:
                        doc.add_picture(tmp_file, width=Inches(4))
                    except:
                        continue
    f = BytesIO()
    doc.save(f)
    f.seek(0)
    return f

def generate_pdf(content, images_dict, title="Assignment"):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", 'B', 16)
    pdf.multi_cell(0, 10, title)
    pdf.ln(10)
    pdf.set_font("Helvetica", '', 12)
    pdf.multi_cell(0, 8, content)
    pdf.ln(5)

    for topic, imgs in images_dict.items():
        if imgs:
            pdf.set_font("Helvetica", 'B', 14)
            pdf.multi_cell(0, 10, f"Images for topic: {topic}")
            for img_url in imgs:
                tmp_file = fetch_and_prepare_image(img_url)
                if tmp_file:
                    try:
                        pdf.image(tmp_file, w=100)
                        pdf.ln(5)
                    except:
                        continue

    f = BytesIO()
    pdf.output(f)
    f.seek(0)
    return f

# ---------------------------- Streamlit UI ----------------------------
assets_path = Path(__file__).parent / "assets"

st.set_page_config(
    page_title="Assignment Generator",
    page_icon=str(assets_path / "logo.png"),
    layout="wide"
)

# 🎉 Fun UI
floating_icons()
render_header(assets_path, lottie_url="https://assets6.lottiefiles.com/packages/lf20_jcikwtux.json")

# Motivational Banner
st.markdown("""
<div style='background: linear-gradient(90deg, #ff6ec4, #42e695);
            padding:20px; border-radius:15px; margin-top:15px; text-align:center; color:white; 
            font-family:Comic Sans MS; box-shadow:0 5px 15px rgba(0,0,0,0.2);'>
    🚀 Let's Explore & Learn! Generate Assignments like a Pro! 📝
</div>
""", unsafe_allow_html=True)

info_box()
st.write("Input your topic keywords and get a structured assignment with images and downloads.")

# ---------------------------- Session State ----------------------------
if "assignment_generated" not in st.session_state:
    st.session_state.assignment_generated = False
    st.session_state.assignment_content = ""
    st.session_state.suggestions = {}
    st.session_state.images_dict = {}

# ---------------------------- Inputs ----------------------------
keywords_input = st.text_area(
    "Enter keywords/topics (comma-separated)", 
    placeholder="e.g., Photosynthesis, Chlorophyll, Light Reactions",
    value=""
)
output_format = st.radio("Select output format:", ["DOCX", "PDF"])
max_points = st.number_input("Maximum key points per topic:", min_value=3, max_value=20, value=10)
max_images = st.number_input("Maximum images per topic:", min_value=1, max_value=5, value=3)

# ---------------------------- Generate Assignment Button ----------------------------
generate_clicked = st.button("📝 Generate Assignment", key="generate_btn")

if generate_clicked or st.session_state.assignment_generated:
    if not st.session_state.assignment_generated:
        if not keywords_input.strip():
            st.warning("Please enter at least one keyword/topic!")
        else:
            keywords = [kw.strip() for kw in keywords_input.split(",") if kw.strip()]
            st.info("Generating assignment... this may take a few seconds.")

            assignment_content, suggestions, images_dict = generate_assignment(
                keywords, max_points=max_points, max_images=max_images
            )

            st.session_state.assignment_content = assignment_content
            st.session_state.suggestions = suggestions
            st.session_state.images_dict = images_dict
            st.session_state.assignment_generated = True

    # Suggestions
    for kw, sug in st.session_state.suggestions.items():
        if sug:
            st.warning(f"Suggestions for '{kw}': {', '.join(sug)}")

    # Assignment Preview
    st.subheader("📄 Assignment Preview")
    st.text_area("Preview", value=st.session_state.assignment_content, height=400)

    # Image Preview + Downloads
    for topic, imgs in st.session_state.images_dict.items():
        if imgs:
            st.subheader(f"Images for topic: {topic}")
            for idx, img_url in enumerate(imgs):
                st.image(img_url, width=300)
                tmp_file = fetch_and_prepare_image(img_url)
                if tmp_file:
                    with open(tmp_file, "rb") as f:
                        st.download_button(
                            label=f"⬇️ Download Image {idx+1} for {topic}",
                            data=f,
                            file_name=f"{topic.replace(' ', '_')}_{idx+1}.jpg",
                            mime="image/jpeg"
                        )

    # Assignment Download
    file_name = "Assignment_" + "_".join([kw.replace(" ", "_") for kw in keywords_input.split(",") if kw.strip()])
    if output_format == "DOCX":
        file_data = generate_docx(st.session_state.assignment_content, st.session_state.images_dict, title=file_name)
        st.download_button(
            "⬇️ Download DOCX",
            data=file_data,
            file_name=f"{file_name}.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )
    else:
        file_data = generate_pdf(st.session_state.assignment_content, st.session_state.images_dict, title=file_name)
        st.download_button(
            "⬇️ Download PDF",
            data=file_data,
            file_name=f"{file_name}.pdf",
            mime="application/pdf"
        )

    st.success("Assignment generated successfully! Images are embedded automatically.")
